"""Extract text from uploaded files (PDF, Word, TXT, Excel, CSV)."""
from typing import Optional
import io


def _decode_bytes(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1256"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def extract_text(uploaded_file) -> Optional[str]:
    """Return the text of an uploaded file, or None if it can't be read."""
    name = (uploaded_file.name or "").lower()
    try:
        uploaded_file.seek(0)
    except Exception:
        pass

    try:
        if name.endswith(".txt"):
            return _decode_bytes(uploaded_file.read())

        if name.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(uploaded_file)
            return "\n".join((page.extract_text() or "") for page in reader.pages)

        if name.endswith(".docx"):
            import docx
            d = docx.Document(uploaded_file)
            parts = [p.text for p in d.paragraphs]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)

        if name.endswith((".xlsx", ".xls", ".csv")):
            return _extract_table(uploaded_file)
    except Exception as e:
        print(f"extract_text error for {name}: {e}")
        return None
    return None


def _extract_table(uploaded_file) -> Optional[str]:
    import pandas as pd
    name = (uploaded_file.name or "").lower()
    if name.endswith(".csv"):
        raw = uploaded_file.read()
        df = pd.read_csv(io.StringIO(_decode_bytes(raw)))
        return df.to_string(index=False)
    sheets = pd.read_excel(uploaded_file, sheet_name=None)  # all sheets
    return "\n\n".join(f"[Sheet: {n}]\n{df.to_string(index=False)}" for n, df in sheets.items())


def get_file_summary(uploaded_file, max_chars: int = 8000) -> dict:
    """Return {name, type, text (truncated), truncated, full_length, error}."""
    name = uploaded_file.name
    text = extract_text(uploaded_file)
    if text is None or not text.strip():
        return {
            "name": name, "type": "unknown", "text": "", "truncated": False,
            "error": "Could not extract any text (unsupported, scanned or empty file)",
        }
    return {
        "name": name,
        "type": name.rsplit(".", 1)[-1].lower() if "." in name else "unknown",
        "text": text[:max_chars],
        "truncated": len(text) > max_chars,
        "full_length": len(text),
        "error": None,
    }
