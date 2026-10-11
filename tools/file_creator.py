"""File creation tool - AI writes the content, we save it as txt/md/docx/pdf."""
import io
import os
import re
from typing import Tuple, Optional
from config import Settings
from tools.chat import run_chat

_ARABIC_RE = re.compile(r"[\u0600-\u06FF]")


def _has_arabic(text: str) -> bool:
    return bool(_ARABIC_RE.search(text))


def _detect_requested_format(prompt: str) -> Optional[str]:
    """Detect an explicitly requested format using WHOLE words (so 'password' != 'word')."""
    p = prompt.lower()
    if re.search(r"(?<![a-z])(docx|word)(?![a-z])", p) or "وورد" in p or "ورد" in p.split():
        return "docx"
    if re.search(r"(?<![a-z])pdf(?![a-z])", p) or "بي دي اف" in p:
        return "pdf"
    if re.search(r"(?<![a-z])(markdown|md)(?![a-z])", p):
        return "md"
    if re.search(r"(?<![a-z])txt(?![a-z])", p) or "نصي" in p:
        return "txt"
    return None


def create_file_from_prompt(user_prompt: str, lang: str = "ar") -> Tuple[Optional[bytes], str, str]:
    """Returns (file_bytes, filename, file_format). Raises LLMError if the LLM call fails."""
    messages = [
        {"role": "system", "content": Settings.SYSTEM_PROMPTS["file_creator"]},
        {"role": "user", "content": user_prompt},
    ]
    reply = run_chat(messages, lang)  # raises LLMError (never turns an error into a file)
    if not reply or not reply.strip():
        return None, "", ""

    file_format = "txt"
    content = reply

    match = re.search(r"\[FILE_TYPE:\s*(\w+)\]", reply, re.IGNORECASE)
    if match:
        file_format = match.group(1).lower().strip()
        content = re.sub(r"\[FILE_TYPE:\s*\w+\]\s*\n?", "", reply, count=1, flags=re.IGNORECASE)

    # remove an accidental wrapping code fence
    content = re.sub(r"^\s*```[a-zA-Z]*\n(.*?)\n```\s*$", r"\1", content, flags=re.DOTALL).strip()

    requested = _detect_requested_format(user_prompt)
    if requested:
        file_format = requested
    if file_format not in ("txt", "md", "docx", "pdf"):
        file_format = "txt"

    # Title: first markdown heading in the content, else the prompt
    title_match = re.search(r"^\s*#{1,3}\s+(.+)$", content, re.MULTILINE)
    title = title_match.group(1).strip() if title_match else user_prompt.strip()

    safe_name = _safe_filename(title)
    filename = f"{safe_name}.{file_format}"

    try:
        if file_format in ("txt", "md"):
            file_bytes = content.encode("utf-8")
        elif file_format == "docx":
            file_bytes = _make_docx(content, title)
        else:
            file_bytes = _make_pdf(content, title, lang)
        return file_bytes, filename, file_format
    except Exception as e:
        print(f"File creation error ({file_format}): {e}")
        return content.encode("utf-8"), f"{safe_name}.txt", "txt"


def _safe_filename(text: str) -> str:
    text = text.strip()[:40]
    text = re.sub(r"[^\w\u0600-\u06FF\s-]", "", text)
    text = re.sub(r"\s+", "_", text).strip("_")
    return text if text else "ai_file"


# ---------------------------------------------------------------- DOCX
def _set_rtl(paragraph):
    """Mark a paragraph as right-to-left (needed for Arabic in Word)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pPr = paragraph._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        pPr.append(OxmlElement("w:bidi"))


def _make_docx(content: str, title: str) -> bytes:
    import docx

    doc = docx.Document()
    heading = doc.add_heading(title[:80], 0)
    if _has_arabic(title):
        _set_rtl(heading)

    skipped_title = False
    for para in content.split("\n"):
        para = para.strip()
        if not para:
            continue
        if not skipped_title and para.lstrip("#").strip() == title.strip() and para.startswith("#"):
            skipped_title = True  # title already rendered above
            continue
        if para.startswith("### "):
            p = doc.add_heading(para[4:], level=3)
        elif para.startswith("## "):
            p = doc.add_heading(para[3:], level=2)
        elif para.startswith("# "):
            p = doc.add_heading(para[2:], level=1)
        elif para.startswith(("- ", "* ")):
            p = doc.add_paragraph(para[2:], style="List Bullet")
        else:
            p = doc.add_paragraph(para)
        if _has_arabic(para):
            _set_rtl(p)

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


# ---------------------------------------------------------------- PDF
_FONT_CANDIDATES = [
    # Windows
    r"C:\Windows\Fonts\arial.ttf",
    r"C:\Windows\Fonts\tahoma.ttf",
    r"C:\Windows\Fonts\segoeui.ttf",
    # Linux
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/noto/NotoNaskhArabic-Regular.ttf",
    "/usr/share/fonts/truetype/amiri/Amiri-Regular.ttf",
    # macOS
    "/Library/Fonts/Arial Unicode.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
]


def _register_font() -> str:
    """Register the first usable TTF that has Arabic glyphs. Falls back to Helvetica."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    local = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "fonts")
    candidates = []
    if os.path.isdir(local):
        candidates += [os.path.join(local, f) for f in sorted(os.listdir(local)) if f.lower().endswith(".ttf")]
    candidates += _FONT_CANDIDATES
    for fp in candidates:
        if os.path.exists(fp):
            try:
                pdfmetrics.registerFont(TTFont("AppFont", fp))
                return "AppFont"
            except Exception:
                continue
    return "Helvetica"


def _shape(text: str) -> str:
    """Reshape + reorder Arabic text so it renders correctly in the PDF."""
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        return get_display(arabic_reshaper.reshape(text))
    except ImportError:
        return text


def _wrap_arabic(text: str, font: str, size: int, max_width: float):
    """Wrap LOGICAL text by measured width, then shape each line separately."""
    from reportlab.pdfbase.pdfmetrics import stringWidth
    words, lines, cur = text.split(), [], ""
    for w in words:
        trial = f"{cur} {w}".strip()
        if stringWidth(_shape(trial), font, size) <= max_width or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return [_shape(l) for l in lines]


def _make_pdf(content: str, title: str, lang: str = "ar") -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    font = _register_font()
    bio = io.BytesIO()
    margin = 50
    doc = SimpleDocTemplate(bio, pagesize=A4, leftMargin=margin, rightMargin=margin,
                            topMargin=margin, bottomMargin=margin)
    avail_width = A4[0] - 2 * margin

    styles = getSampleStyleSheet()
    ltr = ParagraphStyle("Body", parent=styles["Normal"], fontName=font, fontSize=11, leading=16)
    rtl = ParagraphStyle("BodyAr", parent=ltr, alignment=2)  # right aligned
    h_ltr = ParagraphStyle("H", parent=ltr, fontSize=15, leading=20, spaceBefore=8)
    h_rtl = ParagraphStyle("HAr", parent=h_ltr, alignment=2)
    t_ltr = ParagraphStyle("T", parent=styles["Title"], fontName=font, fontSize=20, leading=26)

    def add_line(story, text: str, size: int, base_ltr, base_rtl):
        if _has_arabic(text):
            for line in _wrap_arabic(text, font, size, avail_width):
                story.append(Paragraph(_escape_xml(line), base_rtl))
        else:
            story.append(Paragraph(_escape_xml(text), base_ltr))

    story = []
    add_line(story, title[:80], 20, t_ltr, ParagraphStyle("TAr", parent=t_ltr, alignment=2))
    story.append(Spacer(1, 16))

    skipped_title = False
    for raw in content.split("\n"):
        para = raw.strip()
        if not para:
            story.append(Spacer(1, 8))
            continue
        if not skipped_title and para.startswith("#") and para.lstrip("#").strip() == title.strip():
            skipped_title = True  # title already rendered above
            continue
        is_head = para.startswith("#")
        text = para.lstrip("#").strip() if is_head else (("• " + para[2:]) if para.startswith(("- ", "* ")) else para)
        if is_head:
            add_line(story, text, 15, h_ltr, h_rtl)
        else:
            add_line(story, text, 11, ltr, rtl)

    doc.build(story)
    return bio.getvalue()


def _escape_xml(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
